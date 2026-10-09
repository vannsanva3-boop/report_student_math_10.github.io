import os
import io
from flask import Flask, render_template, request, jsonify, send_file, session, redirect, url_for
from database import get_db, init_db

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as ExcelImage

# កំណត់ឱ្យ Flask ស្គាល់ HTML និង Static Files ដែលនៅខាងក្រៅ Folder
app = Flask(__name__, template_folder='.', static_folder='.', static_url_path='/static')
app.secret_key = "grade10_math_tongpo_secret_2026"

def get_grade(total):
    if total >= 95: return "A"
    if total >= 90: return "B"
    if total >= 80: return "C"
    if total >= 65: return "D"
    if total >= 50: return "E"
    return "F"

@app.before_request
def check_auth():
    # អនុញ្ញាតឱ្យទាញយករូបភាព js css និងទំព័រ login ដោយមិនបាច់ login សិន
    if (request.path.startswith('/static') or 
        request.path in ['/login'] or 
        request.path.endswith(('.png', '.jpg', '.jpeg', '.js', '.css', '.ico'))):
        return None
        
    if 'user' not in session:
        if request.path.startswith('/api/'):
            return jsonify({"error": "Unauthorized"}), 401
        return redirect(url_for('login'))

# បន្ថែម Route នេះដើម្បីឱ្យ Browser ស្គាល់ file main.js និងរូបភាពដែលនៅខាងក្រៅ
@app.route('/<filename>')
def serve_root_files(filename):
    if os.path.exists(filename):
        return send_file(filename)
    return "Not Found", 404

@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "GET":
        if 'user' in session: return redirect(url_for('index'))
        return render_template("login.html")

    username = request.form.get("username", "").strip()
    password = request.form.get("password", "").strip()
    db = get_db()
    user = db.execute("SELECT * FROM users WHERE username = ? AND password = ?", (username, password)).fetchone()
    if user:
        session['user'] = user['username']
        session['fullname'] = user['fullname']
        return redirect(url_for('index'))
    return render_template("login.html", error="ឈ្មោះអ្នកប្រើ ឬពាក្យសម្ងាត់មិនត្រឹមត្រូវ!")

@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for('login'))

@app.route("/")
def index():
    return render_template("index.html")

# ----------------- APIs ----------------- #

@app.get("/api/settings")
def get_settings():
    db = get_db()
    st = db.execute("SELECT * FROM settings WHERE id = 1").fetchone()
    return jsonify(dict(st) if st else {})

@app.post("/api/settings")
def save_settings():
    try:
        db = get_db()
        d = request.json or {}
        db.execute("""
            UPDATE settings SET
                teacher_name = ?,
                teacher_date = ?,
                admin_lunar_date = ?,
                admin_solar_date = ?,
                academic_year = ?
            WHERE id = 1
        """, (
            d.get("teacher_name", "វ៉ាន់ ជីវ៉ា"),
            d.get("teacher_date", "ថ្ងៃទី ៣០ ខែ កញ្ញា ឆ្នាំ២០២៦"),
            d.get("admin_lunar_date", "ថ្ងៃ  ពុធ ៤រោច ខែ  ភទ្របទ  ឆ្នាំ  មមី  អដ្ឋសក័ ព.ស.២៥៧០"),
            d.get("admin_solar_date", "ភ្នំពេញ ត្រូវនឹងថ្ងៃទី  ៣០ ខែ  កញ្ញា  ឆ្នាំ២០២៦"),
            d.get("academic_year", "មុខវិជ្ជាសិក្សា៖ គណិតវិទ្យា /កម្រិតថ្នាក់ ១០ /ក្នុងឆ្នាំ២០២៦")
        ))
        db.commit()
        return jsonify({"success": True})
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.get("/api/dashboard-stats")
def dashboard_stats():
    db = get_db()
    total = db.execute("SELECT COUNT(*) FROM students").fetchone()[0]
    female = db.execute("SELECT COUNT(*) FROM students WHERE gender = 'ស្រី'").fetchone()[0]
    male = db.execute("SELECT COUNT(*) FROM students WHERE gender = 'ប្រុស'").fetchone()[0]
    return jsonify({"total": total, "female": female, "male": male})

@app.get("/api/students")
def get_students():
    db = get_db()
    search = f"%{request.args.get('q', '').strip()}%"

    query = """
        SELECT s.*, 
               COALESCE(sc.behavior, 0) as behavior,
               COALESCE(sc.attendance, 0) as attendance,
               COALESCE(sc.homework, 0) as homework,
               COALESCE(sc.worksheet, 0) as worksheet,
               COALESCE(sc.quiz, 0) as quiz,
               (COALESCE(sc.behavior, 0) + COALESCE(sc.attendance, 0) + COALESCE(sc.homework, 0) + COALESCE(sc.worksheet, 0) + COALESCE(sc.quiz, 0)) as total_score
        FROM students s
        LEFT JOIN math_scores sc ON s.id = sc.student_id
        WHERE (s.name LIKE ? OR s.student_code LIKE ?)
        ORDER BY total_score DESC
    """
    rows = db.execute(query, (search, search)).fetchall()

    students_list = []
    for idx, r in enumerate(rows):
        st = dict(r)
        total = round(st["total_score"], 1)
        avg = round(total / 6.0, 2)

        st["total_score"] = int(total) if total.is_integer() else total
        st["average"] = f"{avg:.2f}"
        st["rank"] = idx + 1
        st["result"] = "ជាប់" if total >= 50 else "ធ្លាក់"
        st["grade"] = get_grade(total)
        students_list.append(st)

    return jsonify({"data": students_list})

@app.post("/api/students/<int:sid>/scores")
def update_scores(sid):
    try:
        db = get_db()
        d = request.json or {}
        db.execute("""
            INSERT INTO math_scores (student_id, behavior, attendance, homework, worksheet, quiz)
            VALUES (?, ?, ?, ?, ?, ?)
            ON CONFLICT(student_id) DO UPDATE SET
                behavior = excluded.behavior,
                attendance = excluded.attendance,
                homework = excluded.homework,
                worksheet = excluded.worksheet,
                quiz = excluded.quiz
        """, (sid, float(d.get("behavior", 0)), float(d.get("attendance", 0)), 
              float(d.get("homework", 0)), float(d.get("worksheet", 0)), float(d.get("quiz", 0))))
        db.commit()
        return jsonify({"success": True})
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.post("/api/students")
def add_student():
    db = get_db()
    name = request.form.get("name", "").strip()
    gender = request.form.get("gender", "ស្រី")
    if not name: return jsonify({"error": "សូមបញ្ចូលឈ្មោះ"}), 400

    last_id = db.execute("SELECT MAX(id) FROM students").fetchone()[0] or 0
    student_code = f"STU-{last_id + 1:02d}"

    cursor = db.cursor()
    cursor.execute("INSERT INTO students (student_code, name, gender) VALUES (?, ?, ?)", (student_code, name, gender))
    new_sid = cursor.lastrowid
    cursor.execute("INSERT INTO math_scores (student_id) VALUES (?)", (new_sid,))
    db.commit()
    return jsonify({"success": True}), 201

@app.delete("/api/students/<int:sid>")
def delete_student(sid):
    db = get_db()
    db.execute("DELETE FROM students WHERE id = ?", (sid,))
    db.commit()
    return jsonify({"success": True})

# ----------------- REAL EXCEL EXPORT (Font Battambang សម្រាប់ឈ្មោះសិស្ស) ----------------- #
@app.get("/export/excel")
def export_real_excel():
    db = get_db()
    settings = db.execute("SELECT * FROM settings WHERE id = 1").fetchone()
    students = db.execute("""
        SELECT s.student_code, s.name, s.gender,
               sc.behavior, sc.attendance, sc.homework, sc.worksheet, sc.quiz,
               (sc.behavior + sc.attendance + sc.homework + sc.worksheet + sc.quiz) as total
        FROM students s
        LEFT JOIN math_scores sc ON s.id = sc.student_id
        ORDER BY total DESC
    """).fetchall()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "លទ្ធផលប្រចាំខែ"
    ws.views.sheetView[0].showGridLines = True

    col_widths = {'A': 4, 'B': 6, 'C': 22, 'D': 8, 'E': 9, 'F': 9, 'G': 11, 'H': 9, 'I': 12, 'J': 10, 'K': 10, 'L': 10, 'M': 8, 'N': 8}
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    font_khmer = Font(name="Khmer OS Siemreap", size=10, color="000000")
    font_bold = Font(name="Khmer OS Siemreap", size=10, bold=True, color="000000")
    font_title_khmer = Font(name="Khmer OS Muol Light", size=11, bold=True, color="000000")
    font_red_bold = Font(name="Khmer OS Siemreap", size=10, bold=True, color="FF0000")
    
    font_battambang_student = Font(name="Battambang", size=10, bold=True, color="000000")

    fill_orange_header = PatternFill(start_color="C55A11", end_color="C55A11", fill_type="solid")
    fill_cream = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    fill_stat_orange = PatternFill(start_color="ED7D31", end_color="ED7D31", fill_type="solid")
    fill_stat_green = PatternFill(start_color="385723", end_color="385723", fill_type="solid")
    fill_stat_blue = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='7F7F7F'),
        right=Side(style='thin', color='7F7F7F'),
        top=Side(style='thin', color='7F7F7F'),
        bottom=Side(style='thin', color='7F7F7F')
    )

    logo_path = "logo.png" if os.path.exists("logo.png") else os.path.join("static", "logo.png")
    if os.path.exists(logo_path):
        try:
            img = ExcelImage(logo_path)
            img.width = 85
            img.height = 85
            ws.add_image(img, 'B2')
        except Exception:
            pass

    ws.merge_cells('E1:K1')
    ws['E1'] = "ព្រះរាជាណាចក្រកម្ពុជា"
    ws['E1'].font = font_title_khmer
    ws['E1'].alignment = Alignment(horizontal='center', vertical='center')

    ws.merge_cells('E2:K2')
    ws['E2'] = "ជាតិ សាសនា ព្រះមហាក្សត្រ"
    ws['E2'].font = font_title_khmer
    ws['E2'].alignment = Alignment(horizontal='center', vertical='center')

    divider_path = "divider.png" if os.path.exists("divider.png") else os.path.join("static", "divider.png")
    if os.path.exists(divider_path):
        try:
            d_img = ExcelImage(divider_path)
            d_img.width = 90
            d_img.height = 22
            ws.add_image(d_img, 'G3')
        except Exception:
            pass

    ws['B5'] = settings['school_name'] if settings else "សាលាបង្វឹក និង បង្រៀនគួរពិសេស តុងប៉"
    ws['B5'].font = font_bold

    ws['B6'] = settings['school_branch'] if settings else "អគារ G   (សាខា ស្វាយប៉ាក)"
    ws['B6'].font = font_bold

    ws.merge_cells('E7:K7')
    ws['E7'] = "របាយការណ៍លទ្ធផលប្រឡងប្រចាំខែ"
    ws['E7'].font = Font(name="Khmer OS Muol Light", size=11, bold=True, color="002060")
    ws['E7'].alignment = Alignment(horizontal='center', vertical='center')

    ws.merge_cells('E8:K8')
    ws['E8'] = settings['academic_year'] if settings else "មុខវិជ្ជាសិក្សា៖ គណិតវិទ្យា /កម្រិតថ្នាក់ ១០ /ក្នុងឆ្នាំ២០២៦"
    ws['E8'].font = Font(name="Khmer OS Siemreap", size=10, bold=True, color="002060", underline="single")
    ws['E8'].alignment = Alignment(horizontal='center', vertical='center')

    ws.merge_cells('B9:D9')
    ws['B9'] = "របាយការណ៍ពិន្ទុគិតជា %"
    for col_idx in range(2, 5):
        cell = ws.cell(row=9, column=col_idx)
        cell.fill = fill_orange_header
        cell.font = Font(name="Khmer OS Siemreap", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal='center', vertical='center')

    percentages = {'E9': "10%", 'F9': "10%", 'G9': "10%", 'H9': "50%", 'I9': "20%", 'J9': "100%"}
    for cid, val in percentages.items():
        ws[cid] = val
        ws[cid].fill = fill_orange_header
        ws[cid].font = Font(name="Khmer OS Siemreap", size=10, bold=True, color="FFFFFF")
        ws[cid].alignment = Alignment(horizontal='center', vertical='center')

    ws.merge_cells('K9:K10')
    ws.merge_cells('L9:L10')
    ws.merge_cells('M9:M10')
    ws.merge_cells('N9:N10')
    ws['K9'] = "មធ្យមភាគ"
    ws['L9'] = "ចំណាត់ថ្នាក់"
    ws['M9'] = "លទ្ធផល"
    ws['N9'] = "និទ្ទេស"

    for c in ['K9', 'L9', 'M9', 'N9']:
        ws[c].fill = fill_cream
        ws[c].font = font_bold
        ws[c].alignment = Alignment(horizontal='center', vertical='center')

    ws['B10'] = "ល.រ"
    ws['C10'] = "ឈ្មោះសិស្ស"
    ws['D10'] = "ភេទ"
    for c in ['B10', 'C10', 'D10']:
        ws[c].fill = fill_cream
        ws[c].font = font_bold
        ws[c].alignment = Alignment(horizontal='center', vertical='center')

    sub_titles = {
        'E10': "ឥរិយាបថ", 'F10': "វត្តមាន", 'G10': "កិច្ចការផ្ទះ",
        'H10': "គណិត", 'I10': "ឆ្លើយសំណួរ", 'J10': "ពិន្ទុសរុប"
    }
    for cid, val in sub_titles.items():
        ws[cid] = val
        ws[cid].fill = fill_orange_header
        ws[cid].font = Font(name="Khmer OS Siemreap", size=10, bold=True, color="FFFFFF")
        ws[cid].alignment = Alignment(horizontal='center', vertical='center')

    for row in ws['B9:N10']:
        for cell in row:
            cell.border = thin_border

    current_row = 11
    for idx, s in enumerate(students):
        tot = round(s["total"], 1)
        avg = round(tot / 6.0, 2)
        res = "ជាប់" if tot >= 50 else "ធ្លាក់"
        g = get_grade(tot)

        ws.cell(row=current_row, column=2, value=idx + 1).alignment = Alignment(horizontal='center')
        
        name_cell = ws.cell(row=current_row, column=3, value=s["name"])
        name_cell.alignment = Alignment(horizontal='left')
        name_cell.font = font_battambang_student

        gender_cell = ws.cell(row=current_row, column=4, value=s["gender"])
        gender_cell.alignment = Align
